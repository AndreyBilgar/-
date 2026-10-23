# Авторское заявление for «Известия вузов. Машиностроение», laid out like the journal's form and filled in from the
# manuscript (title, UDC, authors, organisations, corresponding author). Only signatures are left blank, plus the fields
# the authors have not given: phone, postal address, and identifiers they do not have.
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
HERE=os.path.dirname(os.path.abspath(__file__))
OUT=os.path.join(HERE,'../article/journal/submission/Авторское_заявление.docx')
FONT='Times New Roman'; W=17.0
TITLE='Влияние регулируемых фрикционных тормозов на кинематику легкогазовой пушки'
SECTION='Машиностроение и машиноведение'; SUBSECTION='Расчет и конструирование машин'
UDC='534.1:62-592'
SAR=('Саровский физико-технический институт — филиал федерального государственного автономного образовательного '
     'учреждения высшего образования «Национальный исследовательский ядерный университет «МИФИ» (СарФТИ НИЯУ МИФИ), '
     '607186, Нижегородская обл., г. Саров, ул. Духова, д. 6')
MEPHI=('Федеральное государственное автономное образовательное учреждение высшего образования «Национальный '
       'исследовательский ядерный университет «МИФИ» (НИЯУ МИФИ), 115409, г. Москва, Каширское ш., д. 31')
IPF=('Федеральное государственное бюджетное научное учреждение «Федеральный исследовательский центр Институт прикладной '
     'физики им. А.В. Гапонова-Грехова Российской академии наук» (ИПФ РАН), 603950, г. Нижний Новгород, ул. Ульянова, д. 46')
AUTHORS=[('Балябин Андрей Игоревич',[SAR]),('Воронов Сергей Алексеевич',[MEPHI]),('Герасимов Сергей Иванович',[SAR,IPF]),
         ('Заузолков Михаил Валерьевич',[SAR]),('Ломтева Екатерина Евгеньевна',[SAR])]   # affiliations as in the manuscript header
CONTACT='Балябиным Андреем Игоревичем'; EMAIL='andrewbalyabin@yandex.ru'
IDS=[('SPIN-код РИНЦ','—'),('Scopus Author ID','—'),('ORCID','0009-0005-5576-0183'),('Web of Science Researcher ID','—')]   # corresponding author
DATE='«23» октября 2026 г.'

d=Document()
sec=d.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.left_margin=sec.right_margin=Cm(2.0); sec.top_margin=Cm(1.5); sec.bottom_margin=Cm(1.5)
st=d.styles['Normal']; st.font.name=FONT; st.font.size=Pt(12)
rf=st.element.get_or_add_rPr().find(qn('w:rFonts'))
for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),FONT)
st.paragraph_format.space_after=Pt(0); st.paragraph_format.line_spacing=1.0
def par(align='left',before=0,after=0,tabs=()):
    p=d.add_paragraph(); pf=p.paragraph_format
    pf.alignment={'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY}[align]
    pf.space_before=Pt(before); pf.space_after=Pt(after)
    for pos,kind in tabs: pf.tab_stops.add_tab_stop(Cm(pos),kind)
    return p
def run(p,text,size=None,bold=False,under=False):
    r=p.add_run(text); r.bold=bold; r.underline=under
    if size: r.font.size=Pt(size)
    return r
def field(label,value,before=4):
    """«label ____value____» to the right margin, as on the printed form."""
    p=par(before=before,tabs=[(W,WD_TAB_ALIGNMENT.RIGHT)])
    if label: run(p,label+' ')
    run(p,value+'\t',under=True); return p
def caption(text,align='center',indent=0.0):
    p=par(align=align); p.paragraph_format.left_indent=Cm(indent); run(p,text,size=9.5); return p
# ---------------------------------------------------------------- header
for line in ('В редакцию журнала','«Известия высших учебных заведений. Машиностроение»','105005, Москва, 2-я Бауманская ул., д. 5, стр. 1'):
    run(par('right'),line)
run(par('center',before=14,after=10),'Авторское заявление',size=16,bold=True)
run(par('justify'),'Прошу опубликовать в журнале «Известия высших учебных заведений. Машиностроение» статью')
p=par(before=6,tabs=[(W,WD_TAB_ALIGNMENT.RIGHT)]); run(p,'«%s»'%TITLE+'\t',under=True)
caption('название статьи')
field('В разделе',SECTION,before=2)
field('В подразделе',SUBSECTION)
# ---------------------------------------------------------------- statements of the form
for t in ('1. Мне известно, что я несу всю ответственность за содержание этой статьи и за сам факт ее публикации.',
          '2. Публикация моих материалов осуществляется без выплаты гонорара.',
          '3. Даю разрешение на публикацию статьи в открытом доступе на русском языке, перевод статьи на английский язык и публикацию ее в базах данных (РУНЭБ, Scopus, Web of Science и т.д.).',
          '4. Подтверждаю, что статья не была опубликована в другом издании и в настоящее время не находится на рассмотрении в редакционной коллегии другого издания.'):
    run(par('justify',before=2),t)
p=par(before=2); run(p,'5. ')
for k,(lab,val) in enumerate(IDS):
    run(p,(', ' if k else '')+lab+' '); run(p,' %s '%val,under=True)
p=par(before=2); run(p,'УДК '); run(p,UDC,under=True)
# ---------------------------------------------------------------- authors: everything but the signature
NAME_END=11.8; SIGN_START=13.0
for k,(name,orgs) in enumerate(AUTHORS,1):
    p=par(before=12,tabs=[(NAME_END,WD_TAB_ALIGNMENT.LEFT),(SIGN_START,WD_TAB_ALIGNMENT.LEFT),(W,WD_TAB_ALIGNMENT.RIGHT)]); p.paragraph_format.keep_with_next=True
    run(p,'%d. '%k); run(p,name+'\t',under=True); run(p,'\t'); run(p,'\t',under=True)
    p=par(tabs=[(14.6,WD_TAB_ALIGNMENT.LEFT)]); p.paragraph_format.left_indent=Cm(0.8); p.paragraph_format.keep_with_next=True; run(p,'ФИО автора (полностью)\tподпись',size=9.5)
    p=par('left',before=4,tabs=[(W,WD_TAB_ALIGNMENT.RIGHT)]); p.paragraph_format.keep_with_next=True; p.paragraph_format.keep_together=True; run(p,'; '.join(orgs)+'\t',under=True)
    caption('Полное название организации с указанием почтового индекса и адреса',align='left',indent=0.8)
# ---------------------------------------------------------------- correspondence
field('Всю переписку по поводу публикации прошу вести с',CONTACT,before=14)
p=par(tabs=[(12.0,WD_TAB_ALIGNMENT.LEFT)]); run(p,'\tФИО автора',size=9.5)
field('e-mail',EMAIL,before=4)
field('Контактные телефоны','',before=8)
field('Почтовый адрес с индексом','',before=8)
p=par(before=6,tabs=[(W,WD_TAB_ALIGNMENT.RIGHT)]); run(p,'\t',under=True)
run(par(before=12),DATE)
z=d.settings.element.find(qn('w:zoom'))
if z is not None and z.get(qn('w:percent')) is None: z.set(qn('w:percent'),'100')   # required by the schema
d.save(OUT); print('saved',OUT)
