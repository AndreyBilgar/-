import re, copy
from docx import Document
from docx.shared import Pt, Cm, Emu, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_LINE_SPACING, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import parse_xml, OxmlElement
from omml import omath
FONT='Times New Roman'
class Doc:
    def __init__(s, body_pt=12, line=1.15, indent_cm=1.0, text_width_cm=17.0):
        s.d=Document(); s.fig=0; s.tab=0; s.body_pt=body_pt; s.line=line; s.indent=indent_cm; s.tw=text_width_cm
        sec=s.d.sections[0]
        sec.page_width=Cm(21.0); sec.page_height=Cm(29.7)
        sec.left_margin=Cm(2.5); sec.right_margin=Cm(1.5); sec.top_margin=Cm(2.0); sec.bottom_margin=Cm(2.0)
        sec.footer_distance=Cm(1.0)
        st=s.d.styles['Normal']; st.font.name=FONT; st.font.size=Pt(body_pt)
        rpr=st.element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
        if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),FONT)
        lang=OxmlElement('w:lang'); lang.set(qn('w:val'),'ru-RU'); lang.set(qn('w:eastAsia'),'ru-RU'); rpr.append(lang)
        pf=st.paragraph_format; pf.space_after=Pt(0); pf.space_before=Pt(0); pf.line_spacing=line
        # hyphenation
        settings=s.d.settings.element
        ah=OxmlElement('w:autoHyphenation')
        dts=settings.find(qn('w:defaultTabStop'))
        if dts is not None: dts.addnext(ah)
        else: settings.append(ah)
        z=settings.find(qn('w:zoom'))
        if z is not None and z.get(qn('w:percent')) is None: z.set(qn('w:percent'),'100')
        # footer page number
        fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        fld=parse_xml('<w:fldSimple xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:instr="PAGE"><w:r><w:rPr><w:rFonts w:ascii="%s" w:hAnsi="%s"/><w:sz w:val="20"/></w:rPr><w:t>1</w:t></w:r></w:fldSimple>'%(FONT,FONT))
        fp._p.append(fld)
    # ---------- inline text with $math$ and **bold** and _{sub}/^{sup} not supported in text (use $)
    def _fill(s, p, text, size=None, bold=False, italic=False):
        parts=re.split(r'(\$[^$]+\$)', text)
        for part in parts:
            if not part: continue
            if part.startswith('$') and part.endswith('$') and len(part)>1:
                p._p.append(parse_xml(omath(part[1:-1])))
            else:
                for j,seg in enumerate(re.split(r'(\*\*[^*]+\*\*)', part)):
                    if not seg: continue
                    b=bold
                    if seg.startswith('**') and seg.endswith('**'): seg=seg[2:-2]; b=True
                    r=p.add_run(seg); r.bold=b; r.italic=italic
                    if size: r.font.size=Pt(size)
        return p
    def para(s, text, align='justify', indent=True, size=None, bold=False, italic=False, before=0, after=0, keep_next=False, line=None):
        p=s.d.add_paragraph()
        pf=p.paragraph_format
        pf.alignment={'justify':WD_ALIGN_PARAGRAPH.JUSTIFY,'center':WD_ALIGN_PARAGRAPH.CENTER,'left':WD_ALIGN_PARAGRAPH.LEFT,'right':WD_ALIGN_PARAGRAPH.RIGHT}[align]
        pf.first_line_indent=Cm(s.indent) if indent else Cm(0)
        pf.space_before=Pt(before); pf.space_after=Pt(after)
        if line: pf.line_spacing=line
        if keep_next: pf.keep_with_next=True
        return s._fill(p,text,size,bold,italic)
    def heading(s, text, before=10, after=4):
        p=s.para(text,align='left',indent=False,bold=True,before=before,after=after,keep_next=True)
        return p
    def eq(s, latex, num=None, before=3, after=3):
        p=s.d.add_paragraph(); pf=p.paragraph_format
        pf.first_line_indent=Cm(0); pf.space_before=Pt(before); pf.space_after=Pt(after)
        ts=pf.tab_stops; ts.add_tab_stop(Cm(s.tw/2),WD_TAB_ALIGNMENT.CENTER); ts.add_tab_stop(Cm(s.tw),WD_TAB_ALIGNMENT.RIGHT)
        p.add_run('\t'); p._p.append(parse_xml(omath(latex)))
        if num is not None: p.add_run('\t(%s)'%num)
        return p
    def figure(s, path, width_cm, caption):
        s.fig+=1
        p=s.d.add_paragraph(); pf=p.paragraph_format; pf.alignment=WD_ALIGN_PARAGRAPH.CENTER; pf.first_line_indent=Cm(0)
        pf.space_before=Pt(4); pf.keep_with_next=True; pf.line_spacing=1.0
        p.add_run().add_picture(path, width=Cm(width_cm))
        c=s.para('Рисунок %d – %s'%(s.fig,caption),align='center',indent=False,size=s.body_pt-1,after=6,line=1.0)
        return s.fig
    def table(s, caption, header, rows, widths_cm, size=9.5, header_rows=1, align_first_left=True, merge=None):
        s.tab+=1
        cp=s.para('Таблица %d – %s'%(s.tab,caption),align='left',indent=False,size=s.body_pt-1,before=4,after=2,keep_next=True,line=1.0)
        allrows=header+rows
        t=s.d.add_table(rows=len(allrows), cols=len(widths_cm)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
        t.autofit=False
        for i,row in enumerate(allrows):
            tr=t.rows[i]
            trpr=tr._tr.get_or_add_trPr(); cs=OxmlElement('w:cantSplit'); trpr.append(cs)
            if i<header_rows:
                th=OxmlElement('w:tblHeader'); trpr.append(th)
            for j,val in enumerate(row):
                cell=tr.cells[j]; cell.width=Cm(widths_cm[j]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                p=cell.paragraphs[0]; pf=p.paragraph_format; pf.first_line_indent=Cm(0); pf.line_spacing=1.0
                pf.alignment=WD_ALIGN_PARAGRAPH.LEFT if (j==0 and align_first_left and i>=header_rows) else WD_ALIGN_PARAGRAPH.CENTER
                if i<header_rows-0 and i<header_rows: pass
                s._fill(p,str(val),size=size,bold=(i<header_rows))
                if i<header_rows: pf.keep_with_next=True
        # grid widths
        tbl=t._tbl; grid=tbl.tblGrid
        for j,gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(widths_cm[j]*567)))
        # cell margins
        tblPr=tbl.tblPr; mar=parse_xml('<w:tblCellMar xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:top w:w="15" w:type="dxa"/><w:left w:w="57" w:type="dxa"/><w:bottom w:w="15" w:type="dxa"/><w:right w:w="57" w:type="dxa"/></w:tblCellMar>')
        look=tblPr.find(qn('w:tblLook'))
        if look is not None: look.addprevious(mar)
        else: tblPr.append(mar)
        if merge:
            for (r0,c0,r1,c1) in merge:
                a=t.cell(r0,c0); b=t.cell(r1,c1); txt=a.text; m=a.merge(b)
        sp=s.d.add_paragraph(); sp.paragraph_format.space_after=Pt(2); sp.paragraph_format.line_spacing=0.8
        return t
    def page_break(s):
        s.d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    def save(s, path): s.d.save(path)
